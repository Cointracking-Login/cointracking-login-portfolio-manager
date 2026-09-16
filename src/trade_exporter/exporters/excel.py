from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from trade_exporter.domain.models import ClosedTrade


class ExcelExporter:
    BASE_HEADERS = [
        "дата",
        "тикер",
        "цена входа",
        "цена выхода",
        "финрез",
        "комиссия",
    ]
    EXTRA_HEADERS = ["направление", "количество", "финрез net"]

    def __init__(
        self,
        output_dir: str | Path,
        filename_prefix: str = "trades",
        extra_columns: bool = False,
        include_summary_sheet: bool = True,
    ):
        self.output_dir = Path(output_dir)
        self.filename_prefix = filename_prefix
        self.extra_columns = extra_columns
        self.include_summary_sheet = include_summary_sheet

    def export(
        self,
        trades: list[ClosedTrade],
        broker_name: str,
        account_id: str,
        period_from: datetime,
        period_to_inclusive: datetime,
        display_timezone,
    ) -> Path:
        self.output_dir.mkdir(parents=True, exist_ok=True)
        path = self._new_file_path(
            broker_name,
            account_id,
            period_from,
            period_to_inclusive,
        )

        wb = Workbook()
        ws = wb.active
        ws.title = "Сделки"

        headers = list(self.BASE_HEADERS)
        if self.extra_columns:
            headers += self.EXTRA_HEADERS
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True)

        for trade in trades:
            row = [
                trade.exit_at.astimezone(display_timezone).strftime("%d.%m.%y"),
                trade.ticker,
                float(trade.entry_price),
                float(trade.exit_price),
                float(trade.gross_pnl),
                float(trade.commission),
            ]
            if self.extra_columns:
                row += [
                    trade.direction,
                    trade.quantity,
                    float(trade.net_pnl),
                ]
            ws.append(row)

        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        self._format_sheet(ws)

        if self.include_summary_sheet:
            self._add_summary(wb, trades)

        wb.save(path)
        return path

    def _new_file_path(
        self,
        broker_name: str,
        account_id: str,
        period_from: datetime,
        period_to_inclusive: datetime,
    ) -> Path:
        tail = account_id[-6:] if account_id else "account"
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        filename = (
            f"{self.filename_prefix}_{broker_name}_{tail}_"
            f"{period_from.strftime('%Y%m%d')}_"
            f"{period_to_inclusive.strftime('%Y%m%d')}_{stamp}.xlsx"
        )
        return self.output_dir / filename

    @staticmethod
    def _format_sheet(ws) -> None:
        widths = {1: 12, 2: 16, 3: 16, 4: 16, 5: 14, 6: 14, 7: 14, 8: 12, 9: 14}
        for index, width in widths.items():
            ws.column_dimensions[get_column_letter(index)].width = width

        for row in range(2, ws.max_row + 1):
            ws.cell(row, 3).number_format = "0.######"
            ws.cell(row, 4).number_format = "0.######"
            ws.cell(row, 5).number_format = "0.00"
            ws.cell(row, 6).number_format = "0.00"
            if ws.max_column >= 9:
                ws.cell(row, 9).number_format = "0.00"

    @staticmethod
    def _add_summary(wb: Workbook, trades: list[ClosedTrade]) -> None:
        ws = wb.create_sheet("Итоги")
        gross = sum((t.gross_pnl for t in trades), start=Decimal("0"))
        commission = sum((t.commission for t in trades), start=Decimal("0"))
        net = gross - commission
        winners = sum(1 for t in trades if t.net_pnl > 0)
        losers = sum(1 for t in trades if t.net_pnl < 0)
        total = len(trades)

        rows = [
            ("Закрытых сделок", total),
            ("Прибыльных", winners),
            ("Убыточных", losers),
            ("Win rate", winners / total if total else 0),
            ("Gross P&L", float(gross)),
            ("Комиссия", float(commission)),
            ("Net P&L", float(net)),
        ]
        for key, value in rows:
            ws.append([key, value])

        ws.column_dimensions["A"].width = 24
        ws.column_dimensions["B"].width = 18
        ws["B4"].number_format = "0.00%"
        for cell in ("B5", "B6", "B7"):
            ws[cell].number_format = "0.00"
