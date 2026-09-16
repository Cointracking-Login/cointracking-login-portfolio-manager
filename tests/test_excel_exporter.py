from datetime import datetime, timezone
from decimal import Decimal
from time import sleep

from openpyxl import load_workbook

from trade_exporter.domain.models import ClosedTrade
from trade_exporter.exporters.excel import ExcelExporter

UTC = timezone.utc


def sample_trade():
    return ClosedTrade(
        exit_at=datetime(2026, 9, 11, 12, tzinfo=UTC),
        instrument_id="uid",
        ticker="IMOEXF",
        direction="LONG",
        quantity=1,
        entry_price=Decimal("2300"),
        exit_price=Decimal("2306"),
        gross_pnl=Decimal("60"),
        commission=Decimal("18.30"),
    )


def test_excel_has_exact_base_columns_and_ddmmyy_date(tmp_path):
    exporter = ExcelExporter(tmp_path, extra_columns=False, include_summary_sheet=False)
    path = exporter.export(
        [sample_trade()], "tbank", "123456789", datetime(2026, 9, 11, tzinfo=UTC),
        datetime(2026, 9, 11, 23, 59, tzinfo=UTC), UTC,
    )
    wb = load_workbook(path)
    ws = wb["Сделки"]
    assert [cell.value for cell in ws[1]] == ExcelExporter.BASE_HEADERS
    assert ws["A2"].value == "11.09.26"
    assert ws["B2"].value == "IMOEXF"
    assert ws["E2"].value == 60
    assert ws["F2"].value == 18.3


def test_every_export_gets_a_new_filename(tmp_path):
    exporter = ExcelExporter(tmp_path, include_summary_sheet=False)
    kwargs = dict(
        trades=[sample_trade()],
        broker_name="tbank",
        account_id="123456789",
        period_from=datetime(2026, 9, 11, tzinfo=UTC),
        period_to_inclusive=datetime(2026, 9, 11, 23, 59, tzinfo=UTC),
        display_timezone=UTC,
    )
    first = exporter.export(**kwargs)
    second = exporter.export(**kwargs)
    assert first != second
    assert first.exists()
    assert second.exists()
